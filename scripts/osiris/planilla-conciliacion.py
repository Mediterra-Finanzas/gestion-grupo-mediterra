# Planilla única de conciliación de los 23 contratos de Osiris.
# Lee planilla-datos.json (producción, solo lectura) y escribe un .xlsx local.
# Uso: SP=<carpeta del json> python planilla-conciliacion.py <ruta del xlsx>
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SP = os.environ["SP"]
SALIDA = sys.argv[1]
d = json.load(open(os.path.join(SP, "planilla-datos.json"), encoding="utf-8"))

ARIAL = "Arial"
F = lambda **k: Font(name=ARIAL, size=k.pop("size", 10), **k)
AZUL = "0000FF"
AMARILLO = PatternFill("solid", start_color="FFF2CC")
CABECERA = PatternFill("solid", start_color="1F3864")
GRIS = PatternFill("solid", start_color="F2F2F2")
ROJO_SUAVE = PatternFill("solid", start_color="FCE4E4")
fino = Side(style="thin", color="D9D9D9")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
USD = '#,##0;(#,##0);"-"'
PCT_NUM = '0.0" %"'

wb = Workbook()

def cabecera(ws, fila, titulos, anchos, congelar=True):
    for i, (t, a) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = CABECERA
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDE
        ws.column_dimensions[L(i)].width = a
    ws.row_dimensions[fila].height = 32
    if congelar:
        ws.freeze_panes = ws.cell(row=fila + 1, column=4)

def celda(ws, r, col, v, entrada=False, fmt=None, wrap=False, fill=None, bold=False):
    c = ws.cell(row=r, column=col, value=v)
    c.font = F(color=AZUL, bold=bold) if entrada else F(bold=bold)
    if entrada:
        c.fill = AMARILLO
    elif fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    c.border = BORDE
    return c

# ── Listas (para validaciones) ────────────────────────────────────────────────
ls = wb.active
ls.title = "Listas"
ls["A1"] = "Responsable interno"; ls["B1"] = "Decisión tomada"; ls["C1"] = "Corroboración del pago"; ls["D1"] = "¿Liquida contract fee al obtentor?"
for col in "ABCD":
    ls[col + "1"].font = F(bold=True)
resp = ["Sin asignar"] + d["responsables"]
decisiones = ["Pendiente de revisión", "Mantener registro actual", "Corregir contrato (con documento)",
              "Conservar histórico; contrato es correcto", "Solicitar documento al cliente",
              "Marcar como no aplicable (con justificación)", "Escalar"]
corrob = ["Sin corroborar", "Comprobante de pago adjunto", "Conciliación bancaria", "No aplica"]
liquida = ["Pendiente de revisión", "Sí", "Parcial", "No"]
for i, v in enumerate(resp, start=2): ls.cell(row=i, column=1, value=v).font = F()
for i, v in enumerate(decisiones, start=2): ls.cell(row=i, column=2, value=v).font = F()
for i, v in enumerate(corrob, start=2): ls.cell(row=i, column=3, value=v).font = F()
for i, v in enumerate(liquida, start=2): ls.cell(row=i, column=4, value=v).font = F()
for col, w in zip("ABCD", (34, 44, 30, 32)): ls.column_dimensions[col].width = w
rng_resp = f"Listas!$A$2:$A${len(resp) + 1}"
rng_dec = f"Listas!$B$2:$B${len(decisiones) + 1}"
rng_cor = f"Listas!$C$2:$C${len(corrob) + 1}"
rng_liq = f"Listas!$D$2:$D${len(liquida) + 1}"

# ── Conciliación ──────────────────────────────────────────────────────────────
ws = wb.create_sheet("Conciliación", 0)
ws["A1"] = "Osiris · Conciliación de los 23 contratos"
ws["A1"].font = F(bold=True, size=14)
ws["A2"] = (f"Fuente: producción, solo lectura · fila osiris actualizada {d['filaOsirisActualizada'][:16].replace('T', ' ')} UTC · "
            f"generada {d['generado'][:16].replace('T', ' ')} UTC · importes contractuales, no deuda confirmada ni facturación exigible")
ws["A2"].font = F(italic=True, color="595959")
ws["A3"] = ("Celdas amarillas con texto azul: se completan. El resto sale del sistema y no se edita aquí. "
            "Columnas Y a AP: obtentor validado por contrato y concepto; en contract fee, el neto con fuente única es un impacto calculado pendiente de conciliación, no deuda nueva.")
ws["A3"].font = F(color=AZUL)

titulos = ["N°", "ID contrato", "Cliente", "País", "Firmado", "Concepto", "Tipo / cuotas", "Importe contractual (USD)",
           "Factura", "Pago en el sistema", "Fecha de pago registrada", "Corroboración del pago", "Vencimiento",
           "Fuente canónica", "Registro histórico (Fee Entrada)", "Discrepancia", "Clasificación", "Motivo",
           "Decisión requerida", "Responsable interno", "Decisión tomada", "Documento de respaldo (N° o ruta)",
           "Fecha de revisión", "Observaciones",
           "Obtentor según el cálculo (regla que calza)", "Obtentor según variedades plantadas", "¿Coinciden?",
           "Cobro que usa Pago Obtentores hoy", "Cobro con fuente única",
           "Participación aplicada (%)", "Retención aplicada (%)", "Base del cálculo (USD)",
           "Deuda bruta hoy (USD)", "Neto hoy (USD)",
           "Deuda bruta con fuente única (USD)", "Retención con fuente única (USD)", "Neto con fuente única (USD)",
           "Ítems del cálculo (más de 1 = doble conteo)", "Pagos al obtentor registrados en Osiris",
           "Liquidación previa confirmada al obtentor (USD, neto)", "Documento de la liquidación",
           "Impacto calculado pendiente de conciliación (USD)"]
anchos = [5, 20, 30, 9, 8, 16, 16, 14, 20, 18, 14, 22, 30, 30, 36, 34, 22, 36, 44, 24, 30, 30, 13, 34,
          28, 30, 11, 18, 22, 12, 12, 13, 13, 13, 15, 14, 14, 13, 34, 16, 26, 17]
C_OBT, C_VAR, C_COINC, C_HOY, C_FU, C_PCT, C_WHT, C_BASE, C_BH, C_NH, C_BFU, C_WFU, C_NFU, C_ITEMS, C_POSIRIS, C_LIQ, C_LIQDOC, C_IMPACTO = range(25, 43)
FC = 5
cabecera(ws, FC, titulos, anchos)
for col in range(C_OBT, C_IMPACTO + 1):
    ws.cell(row=FC, column=col).fill = PatternFill("solid", start_color="7F6000")
ws.cell(row=FC, column=C_IMPACTO).comment = Comment(
    "Neto con fuente única menos la liquidación previa confirmada. Es un cálculo sobre cobros sin corroborar: "
    "no es deuda nueva confirmada.", "Planilla")
ws.cell(row=FC, column=C_OBT).comment = Comment(
    "El cálculo asigna el ingreso a todo obtentor cuya regla calce por especie y variedad del contrato. "
    "No mira las variedades plantadas. Validar por contrato y concepto antes de aplicar porcentajes.", "Planilla")
pct_fmt = lambda x: PCT_NUM if isinstance(x, (int, float)) else None
r = FC
for f in d["filas"]:
    r += 1
    conflicto = bool(f["discrepancia"])
    fill = ROJO_SUAVE if conflicto else None
    vals = [f["n"], f["contratoId"], f["cliente"], f["pais"], f["firmado"], f["concepto"], f["tipo"], f["importe"],
            f["factura"], f["pagoSistema"], f["fechaPago"], None, f["vencimiento"], f["fuente"], f["historico"],
            f["discrepancia"], f["clasificacion"], f["motivo"], f["decision"]]
    for col, v in enumerate(vals, start=1):
        if col == 12:
            continue
        celda(ws, r, col, v, fmt=USD if col == 8 else None, wrap=col in (13, 14, 15, 16, 18, 19), fill=fill)
    cor_ini = "No aplica" if f["corroboracion"] == "—" else "Sin corroborar"
    celda(ws, r, 12, cor_ini, entrada=True)
    celda(ws, r, 20, f["responsable"], entrada=True)
    celda(ws, r, 21, "Pendiente de revisión", entrada=True)
    celda(ws, r, 22, None, entrada=True)
    celda(ws, r, 23, None, entrada=True, fmt="yyyy-mm-dd")
    celda(ws, r, 24, None, entrada=True, wrap=True)

    ob = f.get("obtentor")
    if ob:
        alerta_obt = ROJO_SUAVE if ob["coinciden"] == "No" else fill
        celda(ws, r, C_OBT, ob["obtentorCalculo"], wrap=True, fill=alerta_obt)
        celda(ws, r, C_VAR, ob["obtentorVariedades"], wrap=True, fill=alerta_obt)
        celda(ws, r, C_COINC, ob["coinciden"], fill=alerta_obt)
        celda(ws, r, C_PCT, ob["participacion"], fmt=pct_fmt(ob["participacion"]), fill=fill)
        celda(ws, r, C_WHT, ob["retencion"], fmt=pct_fmt(ob["retencion"]), fill=fill)
        if ob["regla"]:
            ws.cell(row=r, column=C_PCT).comment = Comment("Regla: " + ob["regla"], "Planilla")
    if ob and not ob.get("soloValidacion"):
        celda(ws, r, C_HOY, ob["cobroHoy"], fill=fill)
        celda(ws, r, C_FU, ob["cobroFuenteUnica"], wrap=True, fill=fill)
        celda(ws, r, C_BASE, ob["base"], fmt=USD, fill=fill)
        celda(ws, r, C_BH, ob["brutoHoy"], fmt=USD, fill=fill)
        celda(ws, r, C_NH, ob["netoHoy"], fmt=USD, fill=fill)
        celda(ws, r, C_BFU, ob["brutoFU"], fmt=USD, fill=fill)
        celda(ws, r, C_WFU, ob["whtFU"], fmt=USD, fill=fill)
        celda(ws, r, C_NFU, ob["netoFU"], fmt=USD, fill=fill)
        celda(ws, r, C_ITEMS, ob["items"], fmt="0", fill=ROJO_SUAVE if ob["items"] > 1 else fill)
        celda(ws, r, C_POSIRIS, ob["pagosOsiris"], wrap=True, fill=fill)
        celda(ws, r, C_LIQ, None, entrada=True, fmt=USD)
        celda(ws, r, C_LIQDOC, None, entrada=True, wrap=True)
        c = ws.cell(row=r, column=C_IMPACTO, value=f'=IF(ISNUMBER({L(C_NFU)}{r}),{L(C_NFU)}{r}-N({L(C_LIQ)}{r}),"")')
        c.number_format = USD; c.font = F(bold=True); c.border = BORDE; c.alignment = Alignment(vertical="top")
    elif ob:
        celda(ws, r, C_HOY, "—", fill=GRIS)
        celda(ws, r, C_FU, "—", fill=GRIS)
        celda(ws, r, C_BASE, "Sin cálculo de montos en este corte: la unificación de Fee Entrada no cambia royalties", wrap=True, fill=GRIS)
        for col in range(C_BH, C_IMPACTO + 1):
            celda(ws, r, col, None, fill=GRIS)
    else:
        celda(ws, r, C_OBT, "Sin evaluar", fill=GRIS)
        for col in range(C_VAR, C_IMPACTO + 1):
            celda(ws, r, col, None, fill=GRIS)
ultima = r

dv_r = DataValidation(type="list", formula1=rng_resp, allow_blank=False,
                      error="Elige un responsable de la lista (o 'Sin asignar').", errorTitle="Responsable")
dv_d = DataValidation(type="list", formula1=rng_dec, allow_blank=False)
dv_c = DataValidation(type="list", formula1=rng_cor, allow_blank=False)
dv_n = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                      error="Monto neto en USD, mayor o igual a 0.", errorTitle="Liquidación")
for dv in (dv_r, dv_d, dv_c, dv_n): ws.add_data_validation(dv)
dv_c.add(f"L{FC + 1}:L{ultima}")
dv_r.add(f"T{FC + 1}:T{ultima}")
dv_d.add(f"U{FC + 1}:U{ultima}")
dv_n.add(f"{L(C_LIQ)}{FC + 1}:{L(C_LIQ)}{ultima}")
ws.auto_filter.ref = f"A{FC}:{L(C_IMPACTO)}{ultima}"

# ── Totales con fórmulas ─────────────────────────────────────────────────────
t = ultima + 3
ws.cell(row=t, column=6, value="Contract fee · importes contractuales por clasificación").font = F(bold=True)
etq = d["etiquetas"]
clases_fee = [("Pendiente de conciliación", etq["informacion_pendiente"]),
              ("Pendiente confirmado (factura sin pago)", etq["pendiente_confirmado"]),
              ("En conflicto", etq["conflicto"]),
              ("Marcado pagado en el sistema (sin corroborar)", etq["cerrado_con_evidencia"]),
              ("Pago corroborado", etq["pago_corroborado"])]
rango_concepto = f"$F${FC + 1}:$F${ultima}"
rango_clase = f"$Q${FC + 1}:$Q${ultima}"
rango_importe = f"$H${FC + 1}:$H${ultima}"
rango = lambda col: f"${L(col)}${FC + 1}:${L(col)}${ultima}"
filas_tot = []
for i, (rotulo, etiqueta) in enumerate(clases_fee):
    rr = t + 1 + i
    ws.cell(row=rr, column=6, value=rotulo).font = F()
    ws.cell(row=rr, column=7, value=etiqueta).font = F(color="595959")
    c = ws.cell(row=rr, column=8, value=f'=SUMIFS({rango_importe},{rango_concepto},"Contract fee",{rango_clase},G{rr})')
    c.number_format = USD; c.font = F()
    n_ = ws.cell(row=rr, column=9, value=f'=COUNTIFS({rango_concepto},"Contract fee",{rango_clase},G{rr})')
    n_.font = F()
    filas_tot.append(rr)
rs = t + 1 + len(clases_fee)
ws.cell(row=rs, column=6, value="Suma clasificada").font = F(bold=True)
ws.cell(row=rs, column=8, value=f"=SUM(H{filas_tot[0]}:H{filas_tot[-1]})").number_format = USD
ws.cell(row=rs, column=8).font = F(bold=True)
ws.cell(row=rs + 1, column=6, value="Suma bruta de importes de contract fee").font = F()
ws.cell(row=rs + 1, column=8, value=f'=SUMIFS({rango_importe},{rango_concepto},"Contract fee")').number_format = USD
ws.cell(row=rs + 1, column=8).font = F()
ws.cell(row=rs + 2, column=6, value="Cuadre (debe ser 0)").font = F(bold=True)
ws.cell(row=rs + 2, column=8, value=f"=H{rs}-H{rs + 1}").number_format = USD
ws.cell(row=rs + 2, column=8).font = F(bold=True)
ws.cell(row=rs + 3, column=6, value="Filas con 'Sin asignar'").font = F()
ws.cell(row=rs + 3, column=8, value=f'=COUNTIF($T${FC + 1}:$T${ultima},"Sin asignar")').font = F()
ws.cell(row=rs + 4, column=6, value="Filas pendientes de revisión").font = F()
ws.cell(row=rs + 4, column=8, value=f'=COUNTIF($U${FC + 1}:$U${ultima},"Pendiente de revisión")').font = F()
ws.cell(row=rs + 5, column=6, value="Contract fee sin corroborar (filas marcadas pagadas)").font = F()
ws.cell(row=rs + 5, column=8, value=f'=COUNTIFS({rango_concepto},"Contract fee",{rango_clase},"{etq["cerrado_con_evidencia"]}",$L${FC + 1}:$L${ultima},"Sin corroborar")').font = F()
ws.cell(row=rs + 5, column=8).comment = Comment("Baja cuando se marca la corroboración en la columna L con un documento o conciliación.", "Planilla")

# ── Obligación con el obtentor · impacto calculado pendiente de conciliación ──
b = rs + 8
ws.cell(row=b, column=6, value="Obligación con el obtentor por contract fee · impacto calculado pendiente de conciliación (no es deuda nueva confirmada)").font = F(bold=True)
co = d["controlObligacion"]
filas_ob = {}
def linea(clave, rotulo, formula, fmt=USD, bold=False):
    rr = b + 1 + len(filas_ob)
    ws.cell(row=rr, column=6, value=rotulo).font = F(bold=bold)
    c = ws.cell(row=rr, column=8, value=formula)
    c.number_format = fmt; c.font = F(bold=bold)
    filas_ob[clave] = rr
    return rr
sum_cf = lambda col: f'=SUMIFS({rango(col)},{rango_concepto},"Contract fee")'
linea("bh", "Deuda bruta calculada hoy (Pago Obtentores vigente)", sum_cf(C_BH))
linea("nh", "Neto calculado hoy", sum_cf(C_NH))
linea("bfu", "Deuda bruta con fuente única", sum_cf(C_BFU))
linea("wfu", "Retención con fuente única", sum_cf(C_WFU))
linea("nfu", "Neto con fuente única", sum_cf(C_NFU), bold=True)
linea("c1", "Cuadre bruto − retención − neto (debe ser 0)", f"=H{filas_ob['bfu']}-H{filas_ob['wfu']}-H{filas_ob['nfu']}", bold=True)
linea("liq", "Liquidaciones previas confirmadas (columna AN)", sum_cf(C_LIQ))
linea("imp", "Impacto calculado pendiente de conciliación", sum_cf(C_IMPACTO), bold=True)
linea("c2", "Cuadre neto − liquidado − impacto (debe ser 0)", f"=H{filas_ob['nfu']}-H{filas_ob['liq']}-H{filas_ob['imp']}", bold=True)
rr_ctrl = linea("ctrl", "Control: neto de calcularDeudaObtentor sobre todo el blob (valor fijo del sistema)", round(co["netoFU"], 2))
ws.cell(row=rr_ctrl, column=8).font = F(color="595959")
linea("c3", "Diferencia contra la suma por contrato (debe ser 0)", f"=H{filas_ob['ctrl']}-H{filas_ob['nfu']}", bold=True)
linea("dup", "Contratos con más de un ítem en el cálculo (doble conteo)", f'=COUNTIFS({rango_concepto},"Contract fee",{rango(C_ITEMS)},">1")', fmt="0")
linea("obtno", "Contract fee con obtentor del cálculo distinto al de las variedades", f'=COUNTIFS({rango_concepto},"Contract fee",{rango(C_COINC)},"No")', fmt="0")
linea("obtno_todos", "Filas de cualquier concepto con obtentor del cálculo distinto al de las variedades", f'=COUNTIF({rango(C_COINC)},"No")', fmt="0")
linea("obtsd", "Filas sin datos de variedad para comprobar el obtentor", f'=COUNTIF({rango(C_COINC)},"Sin datos")', fmt="0")
ws.cell(row=filas_ob["nfu"], column=9, value="= 390.000 marcados pagados × 70 % − 10 % de retención, si todas las filas usan la misma regla").font = F(italic=True, color="595959")

# ── Cuotas royalty planta ────────────────────────────────────────────────────
wq = wb.create_sheet("Cuotas royalty planta", 1)
wq["A1"] = "Cuotas de royalty planta, una fila por cuota"; wq["A1"].font = F(bold=True, size=12)
wq["A2"] = "Fuente: contratos[].rpPlantaCuotas en producción. No hay monto facturado registrado por cuota."; wq["A2"].font = F(italic=True, color="595959")
tq = ["ID contrato", "Cliente", "ID cuota", "Descripción", "Fecha de evento", "N° plantas", "Factura", "Fecha de pago",
      "Estado en el sistema", "Clasificación", "Motivo", "Responsable interno", "Decisión tomada", "Observaciones"]
cabecera(wq, 4, tq, [20, 30, 20, 22, 13, 11, 14, 13, 16, 26, 22, 24, 30, 34])
rq = 4
for q in d["cuotas"]:
    rq += 1
    for col, v in enumerate([q["contratoId"], q["cliente"], q["cuota"], q["descripcion"], q["fechaEvento"], q["nPlantas"],
                             q["factura"], q["fechaPago"], q["estadoSistema"], q["clasificacion"], q["motivo"]], start=1):
        celda(wq, rq, col, v, fmt="#,##0" if col == 6 else None)
    celda(wq, rq, 12, q["responsable"], entrada=True)
    celda(wq, rq, 13, "Pendiente de revisión", entrada=True)
    celda(wq, rq, 14, None, entrada=True, wrap=True)
if rq > 4:
    dvq_r = DataValidation(type="list", formula1=rng_resp); dvq_d = DataValidation(type="list", formula1=rng_dec)
    wq.add_data_validation(dvq_r); wq.add_data_validation(dvq_d)
    dvq_r.add(f"L5:L{rq}"); dvq_d.add(f"M5:M{rq}")
    wq.auto_filter.ref = f"A4:N{rq}"
wq.cell(row=rq + 2, column=1, value="Cuotas").font = F(bold=True)
wq.cell(row=rq + 2, column=2, value=f"=COUNTA(C5:C{rq})").font = F(bold=True)

# ── Evidencia de los marcados pagados ────────────────────────────────────────
we = wb.create_sheet("Evidencia pagos", 2)
we["A1"] = "Qué sostiene los contract fee marcados pagados y el conflicto"; we["A1"].font = F(bold=True, size=12)
we["A2"] = (f"Bitácora disponible desde {d['inicioBitacora'][:10]}. 'Comprobante' y 'Conciliación' indican si el sistema guarda un documento "
            "de pago o una referencia bancaria: hoy ninguno la tiene, así que ningún pago está corroborado.")
we["A2"].font = F(italic=True, color="595959")
te = ["ID contrato", "Cliente", "Clasificación", "Importe (USD)", "Factura", "Fecha de pago registrada",
      "Comprobante en el sistema", "Conciliación en el sistema", "Eventos en bitácora (campos de pago)",
      "Último cambio de estado de pago", "Origen del estado", "Registro histórico Fee Entrada"]
cabecera(we, 4, te, [20, 30, 26, 13, 14, 14, 13, 13, 12, 50, 34, 34])
re_ = 4
for e in d["evidencia"]:
    re_ += 1
    fill = ROJO_SUAVE if e["historico"] else None
    for col, v in enumerate([e["contratoId"], e["cliente"], e["clasificacion"], e["importe"], e["factura"], e["fechaPago"],
                             e["comprobante"], e["conciliacion"], e["eventos"], e["ultimoCambioPago"], e["origen"], e["historico"]], start=1):
        celda(we, re_, col, v, fmt=USD if col == 4 else None, wrap=col in (10, 11, 12), fill=fill)
we.cell(row=re_ + 2, column=3, value="Total marcado pagado (sin corroborar)").font = F(bold=True)
we.cell(row=re_ + 2, column=4, value=f'=SUMIFS(D5:D{re_},C5:C{re_},"{etq["cerrado_con_evidencia"]}")').number_format = USD
we.cell(row=re_ + 2, column=4).font = F(bold=True)
we.cell(row=re_ + 3, column=3, value="Con fecha de pago").font = F()
we.cell(row=re_ + 3, column=4, value=f'=COUNTIFS(C5:C{re_},"{etq["cerrado_con_evidencia"]}",F5:F{re_},"<>")').font = F()
we.cell(row=re_ + 4, column=3, value="Con comprobante o conciliación").font = F()
we.cell(row=re_ + 4, column=4, value=f'=COUNTIF(G5:G{re_},"Sí")+COUNTIF(H5:H{re_},"Sí")').font = F()

m0 = re_ + 7
we.cell(row=m0, column=1, value="Movimientos que mencionan al obtentor con reglas de contract fee (fuera de Osiris)").font = F(bold=True, size=12)
pob = d.get("pagosObtentorEnOsiris")
we.cell(row=m0 + 1, column=1, value=(
    "La coincidencia es por nombre. No prueba que un movimiento liquide contract fee ni a qué contratos cubre. "
    "La factura 121 de IQ no se asume como liquidación de contract fees: el posible solape y la diferencia de USD 100 "
    "quedan como pendientes documentales. "
    f"Pagos registrados en Osiris (pagosObtentor): {'colección inexistente' if pob is None else str(pob) + ' registros'}. "
    "La ausencia de un registro no demuestra que no se haya liquidado.")).font = F(italic=True, color="595959")
tm = ["Fuente", "Fecha", "Tipo / categoría", "Cuenta", "Moneda", "TC", "Debe", "Haber", "Monto", "Documento", "Glosa",
      "¿Liquida contract fee al obtentor?", "Contratos que cubre (ID)", "Monto neto USD atribuible"]
hm = m0 + 3
for i, tt in enumerate(tm, start=1):
    c = we.cell(row=hm, column=i, value=tt)
    c.font = F(bold=True, color="FFFFFF"); c.fill = CABECERA; c.border = BORDE
    c.alignment = Alignment(wrap_text=True, vertical="center")
we.row_dimensions[hm].height = 32
rm = hm
for m in d.get("movimientosObtentor", []):
    rm += 1
    for col, v in enumerate([m["fuente"], m["fecha"], m["tipo"], m["cuenta"], m["moneda"], m["tc"], m["debe"], m["haber"],
                             m["monto"], m["documento"], m["glosa"]], start=1):
        celda(we, rm, col, v, fmt="#,##0.00" if col in (6, 7, 8, 9) else None, wrap=col in (4, 11))
    celda(we, rm, 12, "Pendiente de revisión", entrada=True)
    celda(we, rm, 13, None, entrada=True, wrap=True)
    celda(we, rm, 14, None, entrada=True, fmt=USD)
if rm > hm:
    dv_l = DataValidation(type="list", formula1=rng_liq, allow_blank=False)
    we.add_data_validation(dv_l); dv_l.add(f"L{hm + 1}:L{rm}")
we.cell(row=rm + 2, column=11, value="Movimientos listados").font = F(bold=True)
we.cell(row=rm + 2, column=12, value=f"=COUNTA(A{hm + 1}:A{rm})" if rm > hm else 0).font = F(bold=True)
we.cell(row=rm + 3, column=11, value="Atribuido como liquidación de contract fee (USD neto)").font = F(bold=True)
we.cell(row=rm + 3, column=14, value=f"=SUM(N{hm + 1}:N{rm})" if rm > hm else 0).number_format = USD
we.cell(row=rm + 4, column=11, value="Debe coincidir con la suma de la columna AN de Conciliación").font = F(italic=True, color="595959")

# ── Leyenda ──────────────────────────────────────────────────────────────────
wl = wb.create_sheet("Leyenda", 0)
lineas = [
    ("Osiris · Planilla única de conciliación", True),
    ("", False),
    ("Qué es", True),
    ("Los 23 contratos de producción, con una fila por concepto (contract fee, royalty planta, royalty comercial).", False),
    ("Las cuotas de royalty planta tienen su propia hoja, una fila por cuota.", False),
    ("Todo sale del sistema en solo lectura. Nada de lo que se escriba aquí modifica el sistema.", False),
    ("", False),
    ("Qué se completa", True),
    ("Solo las celdas amarillas con texto azul: corroboración del pago, responsable interno, decisión tomada, documento, fecha, observaciones "
     "y, en contract fee, la liquidación previa confirmada al obtentor y su documento.", False),
    ("Responsable, decisión y corroboración tienen lista desplegable (hoja Listas).", False),
    ("", False),
    ("Cómo leer la clasificación", True),
    ("Conflicto: dos registros del mismo hecho se contradicen. No se elige ninguno; queda fuera de correos.", False),
    ("Pendiente confirmado: hay número de factura y ningún pago registrado.", False),
    ("Información pendiente: falta un dato. En contract fee sin factura ni pago, el importe es 'pendiente de conciliación'.", False),
    ("Marcado pagado en el sistema: alguien registró factura y pago. NO está corroborado con documento ni conciliación.", False),
    ("Pago corroborado: registro con comprobante o conciliación bancaria. Hoy no hay ninguno.", False),
    ("No aplicable: el contrato no genera ese concepto o todavía no tiene base.", False),
    ("", False),
    ("Fuente canónica", True),
    ("El contract fee se lee del contrato. La fila de Fee Entrada, cuando existe, se muestra como registro histórico y se conserva.", False),
    ("", False),
    ("Obtentor (columnas Y a AP)", True),
    ("En todas las filas se valida el obtentor por contrato y concepto antes de aplicar porcentajes: Y muestra quién calza por la regla "
     "del cálculo, Z quién dicen las variedades plantadas y AA si coinciden. 'No' marca una asignación a revisar; 'Sin datos', un contrato "
     "sin plantaciones que permitan comprobarlo.", False),
    ("Solo en contract fee se calculan montos, con la función real de Pago Obtentores sobre dos entradas: la vigente y la fuente única "
     "(rama detenida). En royalties los montos no se recalculan en este corte.", False),
    ("El neto con fuente única es un IMPACTO CALCULADO PENDIENTE DE CONCILIACIÓN. No es deuda nueva confirmada: se apoya en cobros "
     "marcados pagados que ningún documento corrobora.", False),
    ("Para reconocer un monto, por contrato: corroborar el cobro (L), confirmar la regla aplicada (participación AD y retención AE) y "
     "registrar en AN lo ya liquidado al obtentor por ese contrato, con su documento en AO, para no contarlo dos veces.", False),
    ("La falta de comprobante no demuestra que el cobro no ocurrió. La falta de un registro de pago al obtentor no demuestra que no se liquidó.", False),
    ("AL mayor que 1 significa que el mismo contract fee genera más de una obligación: doble conteo.", False),
    ("La factura 121 de IQ no se asume como liquidación de contract fees. Su posible solape y la diferencia de USD 100 son pendientes documentales.", False),
    ("", False),
    ("Ejemplo de fila completada (formato esperado)", True),
    ("Corroboración: Comprobante de pago adjunto · Responsable: (nombre de la lista) · Decisión: Mantener registro actual · "
     "Documento: TR-2022-0715-BCP · Fecha de revisión: 2026-09-15 · Observaciones: transferencia verificada en cartola", False),
    ("", False),
    ("Controles de la hoja Conciliación", True),
    ("Al pie hay totales por clasificación con fórmulas, el cuadre contra la suma bruta (debe dar 0) y contadores de 'Sin asignar' y pendientes.", False),
    ("Debajo, la obligación con el obtentor: bruto, retención, neto, liquidado e impacto, con dos cuadres y un control contra el cálculo "
     "del sistema sobre todo el blob (las tres diferencias deben dar 0), más los contadores de validación del obtentor.", False),
]
for i, (txt, bold) in enumerate(lineas, start=1):
    c = wl.cell(row=i, column=1, value=txt)
    c.font = F(bold=bold, size=14 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
wl.column_dimensions["A"].width = 130

wb.move_sheet("Listas", offset=len(wb.sheetnames))
wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(SALIDA)
print("escrito " + SALIDA + " · filas=" + str(len(d["filas"])) + " cuotas=" + str(len(d["cuotas"])) +
      " evidencia=" + str(len(d["evidencia"])) + " movimientos=" + str(len(d.get("movimientosObtentor", []))))
